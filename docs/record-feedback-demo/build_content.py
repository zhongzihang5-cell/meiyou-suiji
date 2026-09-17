from pathlib import Path
import json
from statistics import mean
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent


def minute(t):
    h, m = map(int, t.split(':'))
    return h * 60 + m


def duration(n):
    n = round(n)
    h, m = divmod(n, 60)
    return (f'{h}小时' if h else '') + (f'{m}分钟' if m else '') or '0分钟'


def feed_stats(rows):
    times = [minute(t) for t, _ in rows]
    return dict(total=sum(v for _, v in rows), count=len(rows), gap=mean(b-a for a,b in zip(times,times[1:])))


milk_today = [('03:00',120),('06:00',130),('09:00',140),('12:00',150),('15:00',160)]
milk_yesterday = [('03:00',110),('07:00',130),('11:00',130),('14:00',150)]
milk_prior = [600,620,520]
mt, my = feed_stats(milk_today), feed_stats(milk_yesterday)
assert mt['total'] == 700 and my['total'] == 520
assert mt['gap'] == 180 and my['gap'] == 220
assert mean(milk_prior) == 580
milk_source = '今天截至15:00：03:00 / 120毫升，06:00 / 130毫升，09:00 / 140毫升，12:00 / 150毫升，15:00 / 160毫升。昨天截至15:00：03:00 / 110毫升，07:00 / 130毫升，11:00 / 130毫升，14:00 / 150毫升。前3天同期奶量依次为600、620、520毫升。'

nap_today = [('09:00','10:00'),('13:00','14:20')]
nap_yesterday = [('09:10','10:05'),('13:20','14:25')]
nap_sum = lambda rows: sum(minute(b)-minute(a) for a,b in rows)
assert nap_sum(nap_today) == 140 and nap_sum(nap_yesterday) == 120
assert mean([150,150,120]) == 140
night_segments = [(minute('20:30'),1440+minute('00:30')),(1440+minute('00:50'),1440+minute('03:30')),(1440+minute('04:00'),1440+minute('06:30'))]
night_total = sum(b-a for a,b in night_segments)
assert night_total == 550 and max(b-a for a,b in night_segments) == 240
assert sum(night_segments[i+1][0]-night_segments[i][1] for i in range(2)) == 50
diaper_types = ['尿','便','尿','尿','尿便']
assert sum('尿' in t for t in diaper_types) == 4
assert sum('便' in t for t in diaper_types) == 2


def scene(id, label, time, record, receipt, paragraphs=(), detail=(), source='', note='', candidate=False):
    return dict(id=id,label=label,time=time,record=record,receipt=receipt,paragraphs=list(paragraphs),detail=list(detail),source=source,note=note,candidate=candidate)


modules = [
 dict(id='feeding',name='喂奶',symbol='奶',status='首版方向',need='省去翻查和计算，知道今天喝了多少、吃了几顿、间隔有没有变化。',
 conclusion='反馈本次奶量、今天累计和昨日同期差异；奶量变化较大时，补充前3天同期对比，帮助用户回顾变化。',
 outputs=['本次奶量、距上次喂奶的间隔。','今天累计奶量、喂奶次数；与昨天同一时刻相比，多或少多少奶、几顿。','需要进一步回顾变化时，展示前3天同期平均奶量，以及喂奶间隔的变化。'],
 boundary='“变化较大”的触发标准留待编辑调研确认。首版呈现变化，不判断厌奶、变化原因或喂养方案效果。亲喂时长不换算成奶量；混合喂养只汇总已记录的瓶喂量，并明确标注。',
 templates=['已记录喂奶{本次奶量}毫升，距上次{间隔}。今天已记录{次数}顿，共{奶量}毫升；比昨天同期{多/少}{奶量差}毫升、{多/少}{次数差}顿。','按需补充：前3天同期平均{奶量}毫升，今天{多/少}{差值}毫升。今天平均喂奶间隔{时长}，比昨天{延长/缩短}{差值}。'],
 scenes=[
 scene('milk-change','奶量变化','15:00','瓶喂 · 160毫升','已记录喂奶160毫升',[
 '距上次喂奶3小时。',f'今天已记录{mt["count"]}顿，共{mt["total"]}毫升；比昨天15:00时多{mt["total"]-my["total"]}毫升、多1顿。'],[
 '前3天同期平均580毫升，今天多120毫升。','今天平均喂奶间隔3小时，比昨天缩短40分钟。'],milk_source,'本场景手动演示“变化时补充3日回顾”，不代表已确定上线阈值。'),
 scene('milk-mixed','亲喂与瓶喂','15:00','亲喂 · 18分钟','已记录亲喂18分钟',[
 '距上次喂奶2小时30分钟。','今天已记录亲喂3次、瓶喂2次，瓶喂合计260毫升。'],[],
 '今天亲喂记录为06:00、10:00、15:00；瓶喂记录为08:00 / 120毫升、12:30 / 140毫升。本次亲喂15:00开始，记录时长18分钟。','亲喂时长与瓶喂毫升分别展示；不把瓶喂量称作全天总摄入量。'),
 scene('milk-first','首次记录','15:00','瓶喂 · 160毫升','已记录喂奶160毫升',[],[],
 '当前只有本次瓶喂记录。','历史不足时，只确认本次；不把无记录当作昨天没喝奶。')]),

 dict(id='sleep',name='睡眠',symbol='眠',status='首版方向',need='知道今天小睡和昨夜睡眠的次数、时长、时间分布，与宝宝平常是否接近。',
 conclusion='白天反馈小睡累计和入睡、醒来时间的对比；完整夜间反馈总时长、夜醒、夜奶和连续睡眠情况。',
 outputs=['白天：本次时长、今天小睡次数与总时长、上一段清醒时长；与昨天及前3天同期对比。','白天：本次入睡和醒来时间，与前3天同序号小睡的时间对照；有喂奶记录时可展示入睡距上次喂奶多久。','夜间：本次时长、整夜分段数、睡眠总时长、夜醒次数、夜奶次数、最长连续睡眠；与前一晚及前3晚对比。'],
 boundary='夜醒仅在完整夜间记录、分段确由清醒中断时统计，早晨最终醒来不计夜醒；不能把漏记形成的空白算作夜醒。夜奶从喂奶记录单独统计。对比个人历史，不评价睡眠达标，也不推断喂奶导致睡眠变化。',
 templates=['白天：这次睡了{时长}。今天已记录{次数}次小睡，共{总时长}，比昨天同期{多/少}{差值}。本次{入睡时间}入睡、{醒来时间}醒来，{与前3天同序号小睡的客观对比}。','夜间：昨夜共睡{总时长}，分{段数}段，夜醒{次数}次、夜奶{次数}次；比前一晚{时长差、次数差}。最长连续睡了{时长}。'],
 scenes=[
 scene('sleep-nap','白天小睡','14:30','13:00—14:20 · 小睡','已记录睡眠1小时20分钟',[
 '今天已记录2次小睡，共2小时20分钟，比昨天14:30时多20分钟。'],[
 '与前3天同期平均时长相同。','这次13:00入睡、14:20醒来，均在前3天第2次小睡的时间范围内。','入睡前清醒了3小时，距上次喂奶记录40分钟。'],
 '今天09:00—10:00、13:00—14:20；昨天09:10—10:05、13:20—14:25。保存时刻为14:30。前3天同期总时长150、150、120分钟；第2觉入睡时间12:55、13:10、13:20，醒来时间14:10、14:30、14:25。最近一次喂奶开始时间12:20。','同序号小睡对比要求存在可对应的记录；仅展示时间关系。'),
 scene('sleep-night','完整夜间','06:30','04:00—06:30 · 夜间睡眠','已记录睡眠2小时30分钟',[
 '昨夜共睡9小时10分钟，分3段；夜醒2次、夜奶2次。','比前一晚多睡30分钟、少醒1次、少喝1次夜奶。'],[
 '昨夜最长连续睡了4小时。','与前3晚平均相比，多睡10分钟、少醒1次。'],
 '昨夜20:30—00:30、00:50—03:30、04:00—06:30，跨午夜后两段归同一夜；两次清醒间断20和30分钟。喂奶00:35、03:40。前一晚共520分钟、夜醒3次、夜奶3次。前3晚总时长560、540、520分钟，夜醒均3次。示例假设整夜记录完整且两次间断为实际清醒。','保存最后一段后才形成完整昨夜汇总；早晨06:30的最终醒来不计夜醒。'),
 scene('sleep-incomplete','夜间记录不全','06:30','04:00—06:30 · 夜间睡眠','已记录睡眠2小时30分钟',[
 '昨夜已记录2段睡眠，共6小时30分钟。'],[],
 '仅有20:30—00:30、04:00—06:30两段，00:30—04:00的情况未知。','无法确认完整夜间时，不输出夜醒次数、整夜总时长或整夜对比；只汇总已记录部分。')]),

 dict(id='diaper',name='换尿布',symbol='换',status='首版方向',need='省去翻查上次更换时间和一天排便次数，并把喂奶与尿布记录放在一起作初步对照。',
 conclusion='反馈更换间隔、今天累计和有尿、有便的记录次数；有对应数据时，补充奶量与有尿记录的同期变化。',
 outputs=['本次更换时间、距上次更换的间隔。','今天更换总次数、有尿次数、有便次数；尿便都有的一次分别计入两类。','有同时间范围的喂奶记录时，并列奶量变化和有尿的尿布记录变化，供用户结合实际情况判断。'],
 boundary='当前没有精确尿量字段，不增加毫升记录要求。有尿的尿布记录次数不等于实际排尿次数或尿量；类型未记录时只输出更换次数。不推断“奶多所以尿多”，不据此判断更换频率是否合适或红屁屁原因。',
 templates=['已记录换尿布，距上次{间隔}。今天已换{次数}次，其中有尿{次数}次、有便{次数}次。','有依据时补充：截至{时刻}，奶量比昨天同期{多/少}{奶量差}毫升，有尿的尿布记录{多/少}{次数差}次。'],
 scenes=[
 scene('diaper-both','有尿和便便','15:20','换尿布 · 尿便都有','已记录换尿布 · 尿便都有',[
 '距上次更换2小时30分钟。','今天已换5次，其中有尿4次、有便2次。'],[
 '截至15:20，奶量比昨天同期多180毫升，有尿的尿布记录多1次。'],
 '今天05:00尿、08:00便、10:00尿、12:50尿、15:20尿便都有，共5次。昨天截至15:20有尿记录3次。喂奶使用“奶量变化”示例，15:00—15:20无新增喂奶，今日700毫升、昨日520毫升。','尿便都有按一次更换、一次有尿、一次有便统计。并列变化供粗略对照，不换算尿量。'),
 scene('diaper-time','只记录更换','15:20','换尿布','已记录换尿布',[
 '距上次更换2小时30分钟，今天已记录5次更换。'],[],
 '更换时间05:00、08:00、10:00、12:50、15:20；未记录尿便类型。','缺少尿便类型时，不展示有尿、有便的分类次数或奶尿对照。')]),

 dict(id='solids',name='辅食',symbol='食',status='需求待验证',need='曾提出回顾前3天的食物种类和时间，但尚未找到用户记录后必须回查这些信息的具体情境。',
 conclusion='近3天种类与时间对比暂列待验证。当前仅确认新增的辅食记录，保留已有事实，不生成下一步添加建议。',
 outputs=['确认本次食物和记录时间；实际吃下的份量已明确记录时，原样展示并保留单位。','近3天食物种类与时间对比、接受情况回顾，待找到具体回查场景后再决定是否增加。'],
 boundary='不能根据重复食用或未填写反应，输出“排敏通过”“接受良好”或“可以尝试下一种食物”。不将计划添加、准备份量当成实际摄入。已有过敏及反应记录保留原事实。',
 templates=['已记录辅食：{食物名称}，{记录时间}。','实际食用份量已明确时：已记录辅食：{食物名称}，吃了{实际份量及单位}。'],
 scenes=[scene('solids-basic','确认本次记录','12:10','辅食 · 南瓜泥','已记录辅食：南瓜泥，12:10。',[],[],
 '只有本次食物名称和时间；未提供食用份量、接受情况或反应。','此处仅展示保存确认。3日对比不纳入已确认的首版范围；待验证状态只在评审区展示。')]),

 dict(id='supplements',name='营养补剂',symbol='补',status='间隔反馈可做',need='省去翻查补充记录和喂奶时间，用户结合自己的补充安排判断；帮助用户核对是否有遗漏。',
 conclusion='反馈同一补剂距上次补充、距最近一次喂奶的间隔。遗漏核对保留需求方向，基于个人历史的触发规则和准确性待验证。',
 outputs=['本次补剂名称；距上次补充同一补剂的间隔。','距最近一次喂奶的间隔；有明确对照需要时，再补充相关补剂之间的实际间隔。','候选反馈：个人记录规律有可参考性时，呈现近期记录习惯与本次缺少记录的事实，请用户核对是否已补充。'],
 boundary='当前没有补充计划时间字段；历史频率不能等同应服计划，未记录不能认定漏服。系统不判断应补什么、剂量或时间是否正确，不给补服建议。提醒只作为待验证的核对候选。',
 templates=['已记录{补剂名称}。距上次补充该补剂{间隔}，距最近一次喂奶{间隔}。','遗漏核对候选：过去{天数}天，{补剂名称}均记录在{时间范围}；今天截至{时刻}尚未记录，可核对是否已补充。'],
 scenes=[
 scene('supplement-gap','补充间隔','09:00','营养补剂 · 维生素D3','已记录维生素D3',[
 '距上次补充维生素D3 24小时，距最近一次喂奶45分钟。'],[],
 '本次维生素D3为今天09:00，上次为昨天09:00；最近一次喂奶记录为今天08:15。示例时间不代表推荐补充安排。','间隔以记录时间计算，由用户结合已有安排判断。补剂名称需可区分，缺少上次记录则省略对应间隔。'),
 scene('supplement-check','遗漏核对候选','09:00','营养补剂 · 维生素D3','已记录维生素D3',[
 '距上次补充维生素D3 24小时，距最近一次喂奶45分钟。','前3天的铁剂都记录在07:30—07:45，今天截至09:00尚未记录，可核对是否已补充。'],[],
 'D3及喂奶同“补充间隔”场景；另有前3天铁剂记录07:30、07:45、07:40，今天截至09:00无铁剂记录。示例不假设存在应服计划。','待验证的核对提示，手动选中此场景才演示；不代表3天即可建立规律，也不自动判断漏服。',True),
 scene('supplement-first','缺少历史记录','09:00','营养补剂 · 维生素D3','已记录维生素D3',[],[],
 '只有本次D3记录，没有该补剂历史或附近喂奶记录。','不补全历史频率，不生成下一次补充时间。')])
]

common = [
 '反馈顺序为保存确认与本次结果、今天累计、必要的历史对比。默认展示一段核心反馈，更多历史信息展开查看。',
 '所有示例数字均为虚构记录，用于评审文案和交互，不代表喂养标准、推荐作息或补充方案。',
 '今天与昨天及前3天使用相同时刻截止的数据；前3天不含今天。夜间使用完整夜间对比，跨午夜仍归同一夜。',
 '本版演示中的喂奶时间采用开始时间，间隔按开始到开始计算；实际接入需与原记录口径统一。补记时累计按保存时刻计算，缺少相邻记录不强行补算间隔。',
 '所有累计和对比都基于已有记录。无记录不等于未发生；记录范围不足时省略对应对比，不能把无记录日按零纳入基线。',
 '调研形成的是需求与首版方案方向，尚未代表上线效果已经验证。喂奶变化阈值、辅食回查需求、补剂遗漏提示准确性继续验证。'
]

payload=dict(title='宝宝新增记录后的反馈', date='2026年9月11日', modules=modules, common=common)
(OUT/'content.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'data.js').write_text('window.RECORD_FEEDBACK_DEMO = '+json.dumps(payload,ensure_ascii=False,indent=2)+';\n',encoding='utf-8')

md=['# 宝宝新增记录后的反馈调研结论','', '首版优先省去用户查找、汇总和比较历史记录的操作。以下结论保留已讨论的需求边界，并给出可直接评审的反馈文案。','', '所有数字均为演示数据。','']
for m in modules:
    md += [f'## {m["name"]}', '', f'**结论状态：{m["status"]}。** {m["conclusion"]}', '', '**用户需要：**'+m['need'], '', '**可以反馈：**', '']
    md += ['- '+s for s in m['outputs']]
    md += ['', '**文案模板：**', '']+['> '+s for s in m['templates']]
    for s in m['scenes']:
        md += ['', '**'+s['label']+'示例'+('（待验证）' if s['candidate'] else '')+'：**', '', '> '+s['receipt']]
        md += ['> '+p for p in s['paragraphs']]
        if s['detail']: md += ['> 展开回顾：'+' '.join(s['detail'])]
    md += ['', '**输出边界：**'+m['boundary'],'']
md += ['## 共用口径','']+['- '+s for s in common]
md += ['', '## 调研依据', '', '依据用户提供的小红书截图、喂养反馈调研文档及本轮逐项讨论。截图用于识别疑问和行为线索；作者的喂养方法或效果归因不作为产品判断依据。', '', '喂奶：IMG_6524—IMG_6530；睡眠：IMG_6533—IMG_6537、IMG_6541；换尿布：IMG_6543、IMG_6544；辅食：IMG_6545、IMG_6547—IMG_6549；营养补剂：IMG_6551、IMG_6552。', '', '后续验证重点：用户是否少翻历史、少做手工计算，能否正确理解比较对象；遗漏核对的有效性和打扰情况；辅食是否存在具体回查场景。','']
(OUT/'宝宝新增记录反馈调研结论.md').write_text('\n'.join(md),encoding='utf-8')

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.72)
sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    style=doc.styles[name]
    style.font.name='Arial Unicode MS'
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Arial Unicode MS')
    style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.2
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(23)
doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)


def p(text,bold_prefix=None):
    q=doc.add_paragraph()
    if bold_prefix:
        q.add_run(bold_prefix).bold=True
        q.add_run(text)
    else:q.add_run(text)
    return q


def example(s, detail=False):
    doc.add_heading(s['label']+('候选' if s['candidate'] else '')+'文案',2)
    p(s['receipt']+'。' if not s['receipt'].endswith('。') else s['receipt'])
    for t in s['paragraphs']:p(t)
    if detail and s['detail']:p(' '.join(s['detail']), '展开回顾：')


doc.add_paragraph('宝宝新增记录后的反馈调研结论',style='Title')
p('2026年9月11日　｜　喂奶 睡眠 换尿布 辅食 营养补剂')
p('首版优先省去查找、汇总和比较历史记录的操作。喂奶、睡眠、换尿布和补剂间隔可形成反馈方案；辅食对比需求待验证，补剂遗漏提示的触发规则待验证。所有示例数字均为演示数据。')


def section_start(m):
    doc.add_heading(m['name'],1)
    p(m['conclusion'],'调研结论：')
    p(m['need'],'用户需要：')
    p(' '.join(m['outputs']),'反馈内容：')


section_start(modules[0])
example(modules[0]['scenes'][0],True)
p(modules[0]['templates'][0],'通用模板：')
p(modules[0]['boundary'],'输出边界：')
p('已记录亲喂18分钟。距上次喂奶2小时30分钟。今天已记录亲喂3次、瓶喂2次，瓶喂合计260毫升。','混合喂养示例：')

doc.add_page_break()
section_start(modules[1])
example(modules[1]['scenes'][0],True)
example(modules[1]['scenes'][1],True)
p(modules[1]['boundary'],'输出边界：')
p('夜间记录不全时，只说“昨夜已记录2段睡眠，共6小时30分钟”，省略夜醒次数和整夜对比。','信息不足时：')

doc.add_page_break()
section_start(modules[2])
example(modules[2]['scenes'][0],True)
p(modules[2]['boundary'],'输出边界：')
section_start(modules[3])
example(modules[3]['scenes'][0])
p(modules[3]['boundary'],'输出边界：')
p('待找到用户记完后仍需回查的具体情境，再决定是否加入3日种类与时间对比。','后续验证：')

doc.add_page_break()
section_start(modules[4])
example(modules[4]['scenes'][0])
doc.add_heading('遗漏核对候选文案',2)
p(modules[4]['scenes'][1]['paragraphs'][1])
p('此示例用于评审核对提示，不代表已确定触发规则或应服计划。')
p(modules[4]['boundary'],'输出边界：')
doc.add_heading('共用口径',1)
for t in [common[0],common[2],common[3],common[4]]:p(t)
p('喂奶变化阈值由编辑继续调研；辅食回查需求、补剂遗漏提示准确性待验证。产品效果以减少翻查和计算、正确理解反馈及实际核对帮助来验证。','后续验证：')
p('依据原调研附件、用户截图及本轮逐项讨论整理。截图作者的喂养方法和效果归因不作为产品判断依据。','调研依据：')

doc.core_properties.title='宝宝新增记录后的反馈调研结论'
doc.core_properties.subject='五类宝宝记录的反馈内容 文案示例与输出边界'
doc.core_properties.author=''
doc.save(OUT/'宝宝新增记录反馈调研结论.docx')
print(json.dumps(dict(modules=len(modules),scenes=sum(len(m['scenes']) for m in modules),outputs=['宝宝新增记录反馈调研结论.docx','宝宝新增记录反馈调研结论.md','content.json','data.js']),ensure_ascii=False))
